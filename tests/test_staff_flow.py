"""Employee imports and automatic binding in an isolated database; no network."""
import io
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl
from fastapi.testclient import TestClient

from app import bindings, db, employees, feishu, main, staff_template


class StaffFlowTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        for target, attr, value in [(db, "DB_PATH", Path(folder.name) / "staff.db"),
                                    (main, "ADMIN_TOKEN", ""), (feishu, "FEISHU_APP_ID", ""),
                                    (feishu, "FEISHU_APP_SECRET", "")]:
            mock = patch.object(target, attr, value)
            mock.start()
            self.addCleanup(mock.stop)
        network = patch("requests.sessions.Session.request", side_effect=AssertionError("No real network"))
        network.start()
        self.addCleanup(network.stop)
        db.init_db()
        self.client = TestClient(main.app)

    def employee(self, **changes):
        return employees.save_employee({"name": "测试姓名甲", "department": "本地部门", "employee_no": "001", **changes})["id"]

    def remote(self, **changes):
        return {"name": "测试姓名甲", "open_id": "ou_test", "status": {"is_resigned": False}, **changes}

    def row(self, eid):
        return db.query_one("SELECT * FROM employees WHERE id=?", (eid,))

    def test_excel_download_is_blank_seven_columns_and_authenticated(self):
        response = self.client.get("/api/employees/template.xlsx")
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(list(wb.active.values), [staff_template.HEADERS])
        self.assertEqual(wb.active.column_dimensions["C"].number_format, "@")
        self.assertIn("只填写月日", wb.active["E1"].comment.text)
        self.assertIn("完整入职年月日", wb.active["D1"].comment.text)
        self.assertEqual(wb.active.column_dimensions["E"].number_format, "@")
        self.assertEqual(wb.active["G1"].value, "用户 ID（open_id）")
        self.assertEqual(wb.active.column_dimensions["G"].number_format, "@")
        self.assertIn("当前飞书应用", wb.active["G1"].comment.text)
        self.assertEqual(wb.active.auto_filter.ref, "A1:G1")
        self.assertEqual(wb.active.data_validations.dataValidation[0].formula1, '"在职,离职"')
        wb.close()
        with patch.object(main, "ADMIN_TOKEN", "test-secret"):
            self.assertEqual(self.client.get("/api/employees/template.xlsx").status_code, 401)

    def test_excel_roundtrip_preserves_zeroes_dates_status_without_feishu(self):
        wb = openpyxl.load_workbook(io.BytesIO(staff_template.build()))
        wb.active.append(["测试姓名甲", "部门甲", "001", date(2021, 2, 1), date(2000, 2, 29), "在职"])
        wb.active.append(["测试姓名乙", "部门乙", "002", "2023年3月1日", "1990年12月1日", "离职"])
        stream = io.BytesIO()
        wb.save(stream)
        wb.close()
        response = self.client.post("/api/employees/import", files={"file": ("staff.xlsx", stream.getvalue())})
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertTrue(result["ok"])
        self.assertEqual(result["added"], 2)
        self.assertEqual(result["binding"]["pending"], 1)
        self.assertIn("连接飞书", result["binding"]["msg"])
        rows = db.query("SELECT * FROM employees ORDER BY id")
        self.assertEqual((rows[0]["employee_no"], rows[0]["join_date"], rows[0]["birth_date"]), ("001", "2021-02-01", "2000-02-29"))
        self.assertEqual(rows[1]["active"], 0)
        self.assertIsNone(rows[0]["feishu_open_id"])

    def test_excel_open_id_roundtrip_updates_same_employee_and_preserves_blank_id(self):
        eid = self.employee(feishu_open_id="ou_test")
        wb = openpyxl.load_workbook(io.BytesIO(staff_template.build()))
        wb.active.append(["测试姓名甲", "更新部门", "001", "2021-02-01", "02-29", "在职", "ou_test"])
        wb.active.append(["测试姓名乙", "部门乙", "002", "2023-03-01", "12-01", "在职", "ou_other"])
        stream = io.BytesIO()
        wb.save(stream)
        wb.close()
        result = self.client.post("/api/employees/import", files={"file": ("staff.xlsx", stream.getvalue())}).json()
        self.assertEqual((result["added"], result["updated"]), (1, 1))
        self.assertEqual(result["errors"], [])
        self.assertEqual((self.row(eid)["department"], self.row(eid)["feishu_open_id"]), ("更新部门", "ou_test"))
        rows = self.client.get("/api/employees?only_active=false").json()
        self.assertEqual({row["feishu_open_id"] for row in rows}, {"ou_test", "ou_other"})
        self.assertNotEqual(self.row(eid)["identity_status"], "verified")
        raw = "姓名,工号,用户 ID（open_id）\n测试姓名甲,001,\n".encode("utf-8")
        self.assertTrue(self.client.post("/api/employees/import", files={"file": ("staff.csv", raw)}).json()["ok"])
        self.assertEqual(self.row(eid)["feishu_open_id"], "ou_test")

    def test_user_id_csv_aliases_and_export_header(self):
        template = self.client.get("/api/employees/template.csv")
        self.assertIn("用户 ID（open_id）", template.content.decode("utf-8-sig"))
        for index, header in enumerate(("用户ID", "用户 ID", "用户 ID（open_id）", "open_id", "用户ID(open_id)")):
            raw = f"姓名,部门,{header}\n离线导入{index},测试部门,ou_{index}\n".encode("utf-8")
            result = self.client.post("/api/employees/import", files={"file": ("staff.csv", raw)}).json()
            self.assertEqual(result["added"], 1, result)
        self.assertEqual({row["feishu_open_id"] for row in db.query("SELECT * FROM employees")}, {f"ou_{i}" for i in range(5)})

    def test_open_id_format_errors_and_conflicts_do_not_overwrite_employees(self):
        eid = self.employee(feishu_open_id="ou_test")
        before = self.row(eid)
        for value in (123, "user_123", "ou_", "ou_invalid space", "od_department"):
            response = self.client.post("/api/employees", json={"id": eid, "feishu_open_id": value})
            self.assertEqual(response.status_code, 400, response.text)
            self.assertIn("用户 ID", response.json()["detail"])
            self.assertEqual(self.row(eid), before)
        raw = "姓名,部门,工号,用户 ID（open_id）\n另一人,部门乙,002,ou_test\n测试姓名甲,部门乙,001,ou_other\n".encode("utf-8")
        result = self.client.post("/api/employees/import", files={"file": ("staff.csv", raw)}).json()
        self.assertEqual(len(result["errors"]), 2, result)
        self.assertEqual((result["added"], result["updated"]), (0, 0))
        self.assertEqual(self.row(eid), before)

    def test_supplied_id_must_match_live_feishu_name_before_binding(self):
        with patch.object(feishu, "FEISHU_APP_ID", "test-app"), patch.object(feishu, "FEISHU_APP_SECRET", "test-secret"), \
                patch.object(feishu, "list_users", return_value=[self.remote()]):
            response = self.client.post("/api/employees", json={"name": "测试姓名甲", "department": "部门甲", "feishu_open_id": "ou_wrong"}).json()
            self.assertEqual(response["binding"]["matched"], 0)
            self.assertTrue(response["binding"]["errors"])
            self.assertEqual(self.row(response["id"])["feishu_open_id"], "ou_wrong")
            self.assertNotEqual(self.row(response["id"])["identity_status"], "verified")
            corrected = self.client.post("/api/employees", json={"id": response["id"], "feishu_open_id": "ou_test"}).json()
            self.assertEqual(corrected["binding"]["matched"], 1)
            self.assertEqual(self.row(response["id"])["feishu_open_id"], "ou_test")

    def test_unique_name_binds_without_department_or_second_network_read(self):
        eid = self.employee()
        with patch.object(feishu, "get_user") as get_user, patch.object(feishu, "get_department") as get_department:
            result = bindings.auto_bind([eid], users=[self.remote()])
        self.assertEqual(result["matched"], 1)
        self.assertEqual(self.row(eid)["feishu_open_id"], "ou_test")
        self.assertEqual(self.row(eid)["department"], "本地部门")
        get_user.assert_not_called()
        get_department.assert_not_called()

    def test_new_excel_accepts_month_day_and_still_requires_full_join_date(self):
        wb = openpyxl.load_workbook(io.BytesIO(staff_template.build()))
        wb.active.append(['月日员工', '测试部门', 'M01', '2021-09-11', '09-11', '在职'])
        wb.active.append(['缺少年份', '测试部门', 'M02', '09-11', '09-11', '在职'])
        stream = io.BytesIO()
        wb.save(stream)
        wb.close()
        response = self.client.post('/api/employees/import', files={'file': ('staff.xlsx', stream.getvalue())})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['added'], 1)
        self.assertEqual(len(response.json()['errors']), 1)
        rows = self.client.get('/api/employees').json()
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]['birth_date_display'], rows[0]['join_date']), ('09-11', '2021-09-11'))

    def test_remote_same_name_and_conflicting_id_are_not_guessed(self):
        for users in ([self.remote(), self.remote(open_id="ou_other")],
                      [self.remote(), self.remote(name="测试姓名乙")],
                      [self.remote(name="测试姓名甲乙")]):
            with self.subTest(users=users):
                eid = self.employee()
                result = bindings.auto_bind([eid], users=users)
                self.assertEqual(result["matched"], 0)
                self.assertEqual(result["pending"], 1)
                self.assertIsNone(self.row(eid)["feishu_open_id"])
                db.execute("DELETE FROM employees WHERE id=?", (eid,))

    def test_local_same_names_cannot_share_or_split_feishu_identity(self):
        first = self.employee()
        second = self.employee(employee_no="002", department="另一个部门")
        result = bindings.auto_bind(users=[self.remote()])
        self.assertEqual((result["matched"], result["pending"]), (0, 2))
        self.assertIsNone(self.row(first)["feishu_open_id"])
        self.assertIsNone(self.row(second)["feishu_open_id"])

    def test_existing_id_is_not_silently_replaced(self):
        eid = self.employee(feishu_open_id="ou_old")
        result = bindings.auto_bind([eid], users=[self.remote()])
        self.assertEqual(result["matched"], 0)
        self.assertIn("冲突", result["errors"][0]["msg"])
        self.assertEqual(self.row(eid)["feishu_open_id"], "ou_old")

    def test_connection_failure_does_not_undo_employee_save(self):
        with patch.object(feishu, "FEISHU_APP_ID", "test-id"), patch.object(feishu, "FEISHU_APP_SECRET", "test-secret"), \
                patch.object(feishu, "list_users", side_effect=RuntimeError("timeout")):
            response = self.client.post("/api/employees", json={"name": "测试姓名甲", "department": "部门甲", "active": 1})
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertTrue(result["ok"])
        self.assertEqual(result["binding"]["pending"], 1)
        self.assertIsNotNone(self.row(result["id"]))

    def test_save_and_import_invoke_automatic_binding(self):
        with patch.object(feishu, "FEISHU_APP_ID", "test-id"), patch.object(feishu, "FEISHU_APP_SECRET", "test-secret"), \
                patch.object(feishu, "list_users", return_value=[self.remote()]) as directory:
            response = self.client.post("/api/employees", json={"name": "测试姓名甲", "department": "部门甲", "employee_no": "001"})
            self.assertEqual(response.json()["binding"]["matched"], 1)
            raw = "姓名,部门,工号,入职日期,出生年月,在职状态\n测试姓名甲,部门乙,001,2021年2月1日,2000-02-29,在职\n".encode("utf-8")
            imported = self.client.post("/api/employees/import", files={"file": ("staff.csv", raw)}).json()
        self.assertTrue(imported["ok"])
        self.assertEqual(imported["updated"], 1)
        self.assertEqual(imported["binding"]["matched"], 1)
        self.assertEqual(directory.call_count, 2)
        self.assertEqual(self.row(response.json()["id"])["department"], "部门乙")
