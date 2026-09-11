"""Blank employee workbook matching the staff editor's six data fields."""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ("姓名", "部门", "工号", "入职日期", "出生年月", "在职状态")


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "员工资料"
    ws.append(HEADERS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F1"
    for cell, width in zip(ws[1], (20, 28, 20, 22, 22, 18)):
        cell.font = Font(bold=True, color="44405E")
        cell.fill = PatternFill("solid", fgColor="F3F1FA")
        ws.column_dimensions[cell.column_letter].width = width
    ws.column_dimensions["C"].number_format = "@"
    for column in ("D", "E"):
        ws.column_dimensions[column].number_format = "yyyy-mm-dd"
    for address in ("D1", "E1"):
        ws[address].comment = Comment("填写完整日期（年-月-日），例如 2021-02-01；也支持 2021年2月1日。", "填写说明")
    ws["C1"].comment = Comment("按文本填写以保留开头的零。更新现有员工时请保留原工号。", "填写说明")
    ws["F1"].comment = Comment("填写在职或离职。新增时留空默认在职；更新时留空保留原状态。", "填写说明")
    status = DataValidation(type="list", formula1='"在职,离职"', allow_blank=True)
    status.errorTitle, status.error = "在职状态无效", "请选择在职或离职。"
    status.showErrorMessage = True
    ws.add_data_validation(status)
    status.add("F2:F10001")
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
